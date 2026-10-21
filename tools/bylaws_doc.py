# Word + PDF of office/bylaws.json (seal title page, linked TOC). PDF via LibreOffice.
import json,subprocess,asyncio,os
from docx import Document;from docx.shared import Pt,RGBColor,Inches,Cm;from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.oxml.ns import qn;from docx.oxml import OxmlElement
from playwright.async_api import async_playwright
async def seal():
  async with async_playwright() as p:
    b=await p.chromium.launch(executable_path='/usr/bin/google-chrome');pg=await b.new_page(viewport={'width':600,'height':600})
    await pg.set_content(f'<body style="margin:0;background:transparent">{open("office/seal.svg").read().replace("<svg ","<svg width=600 height=600 ",1)}</body>');await pg.screenshot(path='/tmp/seal.png',omit_background=True);await b.close()
asyncio.run(seal())
B=json.load(open('office/bylaws.json'));NAVY=RGBColor(0x1a,0x2a,0x4a);GOLD=RGBColor(0xa8,0x80,0x10)
d=Document();sec=d.sections[0];sec.page_width,sec.page_height=Inches(8.5),Inches(11)
for m in('left_margin','right_margin'):setattr(sec,m,Inches(1))
sec.top_margin=sec.bottom_margin=Inches(.9)
st=d.styles['Normal'];st.font.name='Georgia';st.font.size=Pt(10.5);st.element.rPr.rFonts.set(qn('w:eastAsia'),'Georgia');st.paragraph_format.space_after=Pt(4)
h=d.styles['Heading 1'];h.font.name='Georgia';h.font.size=Pt(15);h.font.color.rgb=NAVY;h.font.bold=True;h.paragraph_format.space_before=Pt(14);h.paragraph_format.space_after=Pt(6);h.paragraph_format.keep_with_next=True
[h.element.rPr.rFonts.attrib.pop(qn('w:'+a),None) for a in('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme')];h.element.rPr.rFonts.set(qn('w:ascii'),'Georgia');h.element.rPr.rFonts.set(qn('w:hAnsi'),'Georgia')
def bm(p,name,i):
  s=OxmlElement('w:bookmarkStart');s.set(qn('w:id'),str(i));s.set(qn('w:name'),name);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(i));p._p.insert(0,s);p._p.append(e)
def link(p,text,anchor):
  hl=OxmlElement('w:hyperlink');hl.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1A2A4A');rp.append(c);r.append(rp);t=OxmlElement('w:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);hl.append(r);p._p.append(hl)
# footer page numbers
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=fp.add_run('SHADYNASTY League Bylaws · Shoebeer Enforcement Division · page ');r.font.size=Pt(8)
for k,v in(('begin',None),('instr',' PAGE '),('end',None)):
  rr=fp.add_run();rr.font.size=Pt(8)
  if v: e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=v
  else: e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),k)
  rr._r.append(e)
sec.different_first_page_header_footer=True
C=lambda t,sz,col=NAVY,b=True,i=False,sp=6:(lambda p:(setattr(p,'alignment',WD_ALIGN_PARAGRAPH.CENTER),setattr(p.paragraph_format,'space_after',Pt(sp)),(lambda r:(setattr(r.font,'size',Pt(sz)),setattr(r.font,'bold',b),setattr(r.font,'italic',i),setattr(r.font.color,'rgb',col)))(p.add_run(t))))(d.add_paragraph())
d.add_paragraph().paragraph_format.space_after=Pt(60)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture('/tmp/seal.png',width=Inches(2.6))
C('SHADYNASTY',34,sp=0);C('LEAGUE BYLAWS',20,GOLD);C('Office of the Commissioner · Shoebeer Enforcement Division',11,RGBColor(0x55,0x55,0x55),False,True,30)
C(B.get('updated_label','Last updated October 2026'),12,NAVY,True);C('A dynasty Superflex league on Sleeper · Est. 2021',10,RGBColor(0x55,0x55,0x55),False)
C('Trash talking and shit-posting are actively encouraged.',10,GOLD,False,True)
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
p=d.add_paragraph('Table of Contents',style='Heading 1')
for i,s in enumerate(B['sections']):
  q=d.add_paragraph();q.paragraph_format.space_after=Pt(1);q.add_run(f'{i}. ' if i else '').font.color.rgb=GOLD;link(q,s['title'],'s_'+s['id'].replace('-','_'))
for i,s in enumerate(B['sections']):
  if i==0: d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
  p=d.add_paragraph((f'§{i}  ' if i else '')+s['title'],style='Heading 1');bm(p,'s_'+s['id'].replace('-','_'),i+1)
  for b in s['b']:
    if b['t']=='table':
      t=d.add_table(rows=0,cols=2);t.style='Light Grid Accent 1'
      for j,row in enumerate(b['rows']):
        cs=t.add_row().cells
        for c,v in zip(cs,row):
          c.text='';r=c.paragraphs[0].add_run(v);r.font.size=Pt(9.5);r.bold=j==0
      for row in t.rows: row.cells[0].width=Inches(2.1);row.cells[1].width=Inches(4.4)
    elif b['t']=='img': d.add_picture(b['src'],width=Inches(5))
    elif b['t']=='li':
      p=d.add_paragraph(b['x'],style='List Bullet' if b['d']<=1 else 'List Bullet 2' if b['d']==2 else 'List Bullet 3')
    else:
      p=d.add_paragraph();r=p.add_run(b['x'])
      if b['x'].endswith(':') or b['x'].isupper() or len(b['x'])<28: r.bold=True;r.font.color.rgb=NAVY
cp=d.core_properties;cp.title='SHADYNASTY League Bylaws (October 2026)';cp.author='SHADYNASTY League Office';cp.comments='';cp.last_modified_by='SHADYNASTY League Office'
out='/workspace/shadynasty_office/SHADYNASTY_Bylaws_2026.docx';d.save(out)
subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir','/workspace/shadynasty_office',out],check=True,capture_output=True)
import shutil;shutil.copy(out,'office/docs/bylaws-2026.docx');shutil.copy(out[:-4]+'pdf','office/docs/bylaws-2026.pdf')

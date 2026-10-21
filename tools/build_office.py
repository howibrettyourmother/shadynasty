# Builds office/ assets from /workspace/shadynasty_office: bylaws.json, seal.svg, cleaned doc HTML (+PNG/PDF via render step)
import docx,json,re,html,os,shutil
SRC='/workspace/shadynasty_office';OUT='office'
d=docx.Document(SRC+'/SHADYNASTY_Bylaws.docx')
TOC=['League Summary','General League Calendar','League Platforms','Scoring','Rosters','Lineups','Tie Breakers','Regular Season, Playoffs, and Toilet Bowl','Fees & Payouts','Trading','Free Agent Acquisition/Waivers','Tanking','Annual Rookie/FA Draft','Startup Draft','Unintentional Roster Moves','Computer Failure','Teams, Ownership, and Replacement','Team Abandonment','Refunds','Anti-competitive Behavior','Activity Rules','Owner Conduct & Responsibilities','Commissioner Responsibilities','League Disbanding','Final Rule','Rule Changes','Appendix']
slug=lambda s:re.sub(r'[^a-z0-9]+','-',s.lower().replace('&','and')).strip('-')
secs=[{'id':'intro','title':'Preamble','b':[]}];cur=secs[0];skip=False
for p in d.paragraphs:
  t=re.sub(r'[\u200b-\u200f\u2060\ufeff\ue000-\uf8ff\x00-\x08\x0b-\x1f]','',p.text.replace('\u00a0',' '));t=re.sub(r'\s+',' ',t).strip()
  if not t: continue
  if t=='SHADYNASTY League Bylaws' or t=='Last updated 1 July 2021': continue
  if t=='Contents': skip=True;continue
  if t in TOC:
    skip=False;cur={'id':slug(t),'title':t,'b':[]};secs.append(cur)
    if t=='General League Calendar':
      tb=d.tables[0];rows=[[re.sub(r'\s+',' ',c.text).strip() for c in r.cells] for r in tb.rows]
      cur['b'].append({'t':'table','rows':[r for r in rows if any(r)]})
    continue
  if skip or cur['id']=='general-league-calendar': continue
  t=re.sub(r'^9\.1\.1\s*','',t)
  n=p._p.pPr.numPr if p._p.pPr is not None else None
  lv=n.ilvl.val if n is not None and n.ilvl is not None else None
  if cur['id']=='league-summary': depth=(lv or 0)+1 if lv is not None else 0
  else: depth=0 if lv is None else max(1,lv-1)
  cur['b'].append({'t':'li' if depth else 'p','d':depth,'x':t})
secs=[s for s in secs if s['b'] or s['id']=='appendix']
os.makedirs(OUT+'/img',exist_ok=True)
import zipfile;zipfile.ZipFile(SRC+'/SHADYNASTY_Bylaws.docx').extract('word/media/image1.gif','/tmp/bx');shutil.copy('/tmp/bx/word/media/image1.gif',OUT+'/img/bracket-example.gif')
for s in secs:
  if s['id']=='appendix': s['b']=[{'t':'p','d':0,'x':'Playoff Bracket Example:'},{'t':'img','src':'office/img/bracket-example.gif','alt':'Example six-team playoff bracket with two first-round byes'}]
json.dump({'updated':'2021-07-01','sections':secs},open(OUT+'/bylaws.json','w'),ensure_ascii=False,indent=0)
L=open(SRC+'/l.html').read();M=open(SRC+'/m.html').read()
svg=re.search(r'<svg.*?</svg>',L,re.S).group(0)
seal=re.sub(r'<svg[^>]*>','<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100" role="img" aria-label="Shoebeer Enforcement Division seal">',svg,1)
open(OUT+'/seal.svg','w').write(seal+'\n')
ADDR=re.compile(r'\b\d{1,6}\s+(?:[A-Z][a-z]+\s+){1,3}(?:Rd|Road|St|Street|Ave|Avenue|Dr|Drive|Ln|Lane|Way|Blvd|Ct|Court|Pl|Place|Pkwy|Cir|Ter|Hwy)\b\.?|Randolph',re.I)
def clean(h,title):
  h=ADDR.sub('',h);return h.replace('<html><head>',f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{title}</title>',1)
open(OUT+'/docs/shady-2026-001-01.html','w').write(clean(L,'SHADYNASTY League Office · Notice SHADY-2026-001.01'))
open(OUT+'/docs/memo-2026-01.html','w').write(clean(M,'SHADYNASTY League Office · Memo 2026-01'))
print([ (s['id'],len(s['b'])) for s in secs])
